"""내 기록 보기 — PC의 매크로 기록 워크북을 읽어 화면용 JSON을 만든다(읽기 전용).

브라우저 작업 스레드(web/worker.mjs)가 사용자가 고른 파일을 메모리에 올린 뒤 이 모듈을 부른다.
파일은 어디에도 저장하거나 보내지 않으며, 이 모듈도 워크북에 쓰지 않는다.
점수 합계·판정은 engine/sheet_v2.read_state가 다시 계산한 값을 쓴다(엑셀 수식 계산값에 기대지 않는다).
News DB 행을 고르는 규칙(5행부터, '예시' 행·비표준 요인 제외)은 engine/rollup_analysis와 같다.
"""
import datetime as dt
import json
import zipfile

import openpyxl

import sheet_v2 as sv
from rollup_analysis import FACTORS, parse_date, week_bounds

NEWS_FIELDS = ("date", "factor", "region", "keyword", "title", "content", "direction", "intensity", "source", "kind")


def _text(value):
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value).strip()


def _number(value):
    try:
        return None if value is None or value == "" else float(value)
    except (TypeError, ValueError):
        return None


def _news(ws):
    out = []
    for index, row in enumerate(ws.iter_rows(min_row=5, max_col=10, values_only=True), start=5):
        row = tuple(row) + (None,) * (10 - len(row))
        if str(row[0]).strip() == "예시" or row[1] not in FACTORS:
            continue
        item = dict(zip(NEWS_FIELDS, row))
        day = parse_date(item["date"])
        out.append({**{key: _text(value) for key, value in item.items()},
                    "date": day.isoformat() if day else _text(item["date"]),
                    "intensity": _number(item["intensity"]),
                    "sheet_row": index})
    # 최신 날짜가 위로. 같은 날은 시트에서 아래(나중에 추가된) 행이 위로. 날짜가 없으면 맨 아래.
    return sorted(out, key=lambda n: (n["date"], n["sheet_row"]), reverse=True)


def read(path):
    state = sv.read_state(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Analysis"]
        week = parse_date(ws["I1"].value)
        conclusion = _text(ws["D3"].value)
        digest = {r: _text(ws.cell(r, sv.COL["digest"]).value) for r in sv.ROWS}
        memo = [{"label": _text(ws.cell(r, 2).value), "value": _text(ws.cell(r, 3).value)} for r in sv.MEMO_ROWS]
        news = _news(wb["News DB"])
    finally:
        wb.close()
    for row in state["rows"]:
        row["digest"] = digest[row["row"]]
        for key in ("note", "dur_tag", "cur_tag", "house"):
            row[key] = _text(row[key])
    start, end = week_bounds(week) if week else (None, None)
    return {
        "week": {"date": week and week.isoformat(), "start": start and start.isoformat(), "end": end and end.isoformat()},
        "conclusion": conclusion,
        "state": state,
        "divergences": sv.divergences(state),
        "memo": [m for m in memo if m["label"] or m["value"]],
        "news": news,
    }


def read_json(path):
    """(성공 여부, JSON 문자열). 실패하면 화면에 보여 줄 문구를 담는다."""
    def failure(message):
        return False, json.dumps({"error": message}, ensure_ascii=False)
    try:
        return True, json.dumps(read(path), ensure_ascii=False)
    except sv.SchemaError:
        return failure("v2 형식 기록 파일이 아닙니다. Analysis 시트 W1 칸이 'schema=v2'인 워크북을 여세요.")
    except KeyError:
        return failure("'Analysis'와 'News DB' 시트가 있는 기록 워크북이 아닙니다.")
    except (zipfile.BadZipFile, OSError, ValueError):
        return failure("엑셀 파일을 읽지 못했습니다. 엑셀에서 저장하는 중이었다면 잠시 뒤 다시 읽습니다.")
    except Exception:
        return failure("기록 파일을 읽는 중 문제가 생겼습니다. 파일을 다시 열어 주세요.")
