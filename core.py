"""Offline demonstration. Every write is confined to the supplied demo storage.

The engine directory is a documented snapshot of the existing project's rules.
There is no news classifier or LLM call in this demo: fixtures are explicit.
"""
from __future__ import annotations

import copy
from contextlib import contextmanager
import datetime as dt
import hashlib
import json
import math
from pathlib import Path
import re
import sqlite3
import sys
import time
import uuid

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE / "engine"))
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
import append_news
import draft_apply
import migrate_v2
import rollup_analysis
import sheet_v2 as sv
import update_analysis
from safe_io import safe_save


def json_write(path, data):
    path = Path(path)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)


def sample():
    return json.loads((BASE / "data/sample.json").read_text(encoding="utf-8"))


def validate_item(item):
    """Demo ingress rule, stricter than the legacy append warning-only checks."""
    if not isinstance(item, dict):
        raise ValueError("뉴스는 JSON 객체여야 합니다.")
    for key, choices in (("factor", append_news.FACTORS), ("region", append_news.REGIONS),
                         ("direction", append_news.DIRECTIONS), ("kind", append_news.KINDS)):
        if item.get(key) not in choices:
            raise ValueError(f"{key}: 허용된 값이 아닙니다.")
    dt.date.fromisoformat(item["date"])
    value = item.get("intensity")
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError("강도는 유한한 숫자여야 합니다.")
    direction = item["direction"]
    if ((direction == "금리상방" and value <= 0) or
        (direction == "금리하방" and value >= 0) or
        (direction == "중립" and value != 0)):
        raise ValueError("방향성과 강도 부호가 일치하지 않습니다.")
    return True


def make_workbook(path):
    """Create a synthetic workbook from scratch, never copy personal workbook data."""
    wb = openpyxl.Workbook()
    nd = wb.active
    nd.title = "News DB"
    ws = wb.create_sheet("Analysis")
    nd["A1"] = "가상 예시 데이터 · 채용 포트폴리오 시연 전용"
    headers = ["날짜", "요인", "지역", "핵심키워드", "제목", "핵심내용", "방향성", "강도", "출처/링크", "유형"]
    for col, header in enumerate(headers, 1):
        nd.cell(4, col, header)
    # The sentinel is required by the legacy append template and excluded from aggregation.
    nd["A5"], nd["E5"] = "예시", "가상 데모 서식 행 · 집계 제외"
    table = Table(displayName="NewsDB", ref="A4:J5")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    nd.add_table(table)
    ws["B1"] = "채권 리서치 검토 · 가상 예시"
    ws["D3"] = "【금주 결론】 시연용 초안 — 실제 투자 판단이 아님"
    for r in sv.ROWS:
        ws.cell(r, 2, sv.ROW_LABEL[r])
    # The existing schema builder handles weights, formulas and initial flags.
    migrate_v2.migrate_rows(ws)
    migrate_v2.add_columns(ws)
    migrate_v2.add_flags_and_memo(ws)
    migrate_v2.add_draft_layout(ws)
    migrate_v2.apply_compact_layout(ws)
    ws[sv.SCHEMA_CELL] = sv.SCHEMA_V2
    for sheet in (nd, ws):
        sheet.freeze_panes = "D8" if sheet == ws else "A5"
        for row in sheet:
            for cell in row:
                cell.font = Font(name="맑은 고딕", size=11, color="173247")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.sheet_view.showGridLines = False
    for cell in nd[4]:
        cell.fill = PatternFill("solid", fgColor="173247")
        cell.font = Font(name="맑은 고딕", bold=True, color="FFFFFF")
    for letter, width in {"A":15,"B":17,"C":12,"D":20,"E":35,"F":65,"G":15,"H":10,"I":22,"J":12}.items():
        nd.column_dimensions[letter].width = width
    ws.column_dimensions["D"].width = 72
    safe_save(wb, str(path), expect_sheets=["News DB", "Analysis"], do_backup=False)
    wb.close()


class DemoStore:
    def __init__(self, storage):
        self.root = Path(storage).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self.db = self.root / "history.sqlite3"
        with self.connect() as con:
            con.execute("CREATE TABLE IF NOT EXISTS runs (id TEXT PRIMARY KEY, created TEXT NOT NULL, scenario TEXT NOT NULL, status TEXT NOT NULL, elapsed REAL, error TEXT)")

    @contextmanager
    def connect(self):
        con = sqlite3.connect(self.db)
        con.row_factory = sqlite3.Row
        try:
            with con:
                yield con
        finally:
            con.close()

    def directory(self, run_id):
        if not isinstance(run_id, str) or not re.fullmatch(r"[a-f0-9]{32}", run_id):
            raise ValueError("실행 ID가 올바르지 않습니다.")
        with self.connect() as con:
            if not con.execute("SELECT 1 FROM runs WHERE id=?", (run_id,)).fetchone():
                raise ValueError("실행을 먼저 생성하세요.")
        return self.root / run_id

    def trace(self, directory, stage, **fields):
        entry = {"time":dt.datetime.now(dt.timezone.utc).isoformat(), "stage":stage, **fields}
        with (directory / "trace.jsonl").open("a", encoding="utf-8") as file:
            file.write(json.dumps(entry, ensure_ascii=False) + "\n")

    def append_once(self, directory, items):
        """Whole-batch dedup for this serial local demo, not semantic news dedup."""
        digest = hashlib.sha256(json.dumps(items, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
        marker = directory / "input_hash.txt"
        if marker.exists() and marker.read_text() == digest:
            self.trace(directory, "append", result="duplicate_skipped", count=0)
            return False
        for item in items:
            validate_item(item)
        json_write(directory / "items.json", {"items":items})
        append_news.append(str(directory / "demo.xlsx"), str(directory / "items.json"))
        marker.write_text(digest)
        self.trace(directory, "append", result="ok", count=len(items))
        return True

    def run(self, scenario="baseline"):
        if scenario not in {"baseline", "coerce", "policy_gap"}:
            raise ValueError("알 수 없는 시나리오입니다.")
        run_id = uuid.uuid4().hex
        directory = self.root / run_id
        directory.mkdir()
        start = time.perf_counter()
        with self.connect() as con:
            con.execute("INSERT INTO runs VALUES (?,?,?,?,?,?)", (run_id, dt.datetime.now().isoformat(timespec="seconds"), scenario, "running", None, None))
        try:
            data = sample()
            draft = copy.deepcopy(data["draft"])
            if scenario == "coerce":
                draft["rows"][0].update(dur=0.7, cur=-3)
            elif scenario == "policy_gap":
                draft["rows"][2].update(market_pct=5.1, my_path_pct=4.5)
            self.trace(directory, "start", scenario=scenario, mode="offline_fixture", llm_called=False)
            make_workbook(directory / "demo.xlsx")
            self.append_once(directory, data["items"])
            auto = rollup_analysis.aggregate(data["items"])
            auto["date"] = data["date"]
            json_write(directory / "aggregate.json", auto)
            update_analysis.apply(str(directory / "demo.xlsx"), auto)
            self.trace(directory, "aggregate", opinions_excluded=auto["opinion_n"])
            json_write(directory / "draft.json", draft)
            self.apply(directory, draft)
            duration = round(time.perf_counter() - start, 3)
            with self.connect() as con:
                con.execute("UPDATE runs SET status='ok', elapsed=? WHERE id=?", (duration, run_id))
            self.trace(directory, "complete", seconds=duration)
        except Exception as exc:
            with self.connect() as con:
                con.execute("UPDATE runs SET status='failed', error=? WHERE id=?", (str(exc),run_id))
            self.trace(directory, "failed", error=str(exc))
            raise
        return self.state(run_id)

    def apply(self, directory, draft):
        result = draft_apply.apply_draft(str(directory / "demo.xlsx"), draft)
        json_write(directory / "last_apply.json", result)
        state = sv.read_state(str(directory / "demo.xlsx"))
        missing = [r for r in draft["rows"] if r["row"] in (10,11,12) and
                   not (("market_pct" in r and "my_path_pct" in r) or ("tp_4w_change_bp" in r and "sigma_bp" in r))]
        conclusion = f"【금주 결론】 듀레이션 {state['dur']['verdict']}({state['dur']['total']:+.2f}) · 커브 {state['cur']['verdict']}({state['cur']['total']:+.2f}) — 초안\n• 가상 예시 데이터 기반 검토용 산출물\n• 정책·TP 수치 미입력 {len(missing)}행 — 0점 처리로 종합 점수 희석 가능"
        auto = json.loads((directory / "aggregate.json").read_text(encoding="utf-8"))
        auto["conclusion"] = conclusion
        update_analysis.apply(str(directory / "demo.xlsx"), auto)
        self.trace(directory, "draft_apply", **result)

    def review(self, run_id, score):
        if isinstance(score, bool) or not isinstance(score, (int,float)) or not math.isfinite(score) or score not in [n/2 for n in range(-4,5)]:
            raise ValueError("점수는 −2~2 사이의 0.5 단위 숫자여야 합니다.")
        directory = self.directory(run_id)
        wb = openpyxl.load_workbook(directory / "demo.xlsx")
        # Explicit UI reviewer action on a synthetic workbook, not automated draft writing.
        before = wb["Analysis"]["G8"].value
        wb["Analysis"]["G8"] = score
        safe_save(wb, str(directory / "demo.xlsx"), expect_sheets=["News DB","Analysis"])
        wb.close()
        self.trace(directory, "human_review", cell="G8", before=before, after=score,
                   note="실제 사용자 버튼 입력. 같은 값이면 엔진의 변경 감지 대상이 아님.")
        return self.state(run_id)

    def reapply(self, run_id):
        directory = self.directory(run_id)
        self.append_once(directory, sample()["items"])
        draft = json.loads((directory / "draft.json").read_text(encoding="utf-8"))
        self.apply(directory, draft)
        return self.state(run_id)

    def state(self, run_id):
        directory = self.directory(run_id)
        with self.connect() as con:
            run = dict(con.execute("SELECT * FROM runs WHERE id=?", (run_id,)).fetchone())
        if run["status"] != "ok":
            return {"run":run}
        draft = json.loads((directory / "draft.json").read_text(encoding="utf-8"))
        normalized, _ = draft_apply.normalize(draft)
        workbook = directory / "demo.xlsx"
        state = sv.read_state(str(workbook))
        for row in state["rows"]:
            row["proposed_dur"] = normalized[row["row"]]["dur"]
            row["missing_basis"] = row["row"] in (10,11,12) and ("산출 불가" in (row["note"] or "") or "미입력" in (row["note"] or ""))
        return {"run":run, "state":state, "input":sample(), "draft":draft,
                "apply":json.loads((directory / "last_apply.json").read_text(encoding="utf-8")),
                "news_count":len(rollup_analysis.read_week_rows(str(workbook),dt.date(2026,9,21))),
                "trace":[json.loads(line) for line in (directory / "trace.jsonl").read_text(encoding="utf-8").splitlines()]}

    def history(self):
        with self.connect() as con:
            return [dict(row) for row in con.execute("SELECT * FROM runs ORDER BY created DESC, rowid DESC LIMIT 40")]
