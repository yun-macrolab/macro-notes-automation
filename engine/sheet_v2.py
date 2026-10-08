#!/usr/bin/env python3
"""채권 주간 평가 시트 v2 — 레이아웃·규약의 단일 출처.

핵심 원칙: **자동화는 초안을 쓰고, 사람이 조정해 확정한다. 사람이 고친 칸은 다시 덮지 않는다.**
  초안 칸   D(요약·근거) · G/J(금주 점수) · M/N(태그)  → draft_apply.py만, 보호 규칙을 거쳐 기입
  사람 전용 FLAG(C19:C24) · 채점 확정(C26) · 메모       → 자동화 쓰기 금지
  자동 칸   O/P(뉴스 참고 점수) · Q(뉴스 다이제스트) · R(하우스 뷰) · F/I(전주 이월) · D3 · I1
  장부      X~AB(자동화가 마지막으로 쓴 값) · AC(사람이 고친 칸) — 숨김

행(8~14): 8 펀더멘털·글로벌 / 9 펀더멘털·국내 / 10 통화정책·글로벌 / 11 통화정책·국내
          12 수급·글로벌 기간 프리미엄 / 13 수급·국내 / 14 기술적·포지션
부호: 듀레이션 + = 금리 하락 요인, 커브 + = 플래트닝. (News DB 강도는 반대: + = 금리상방)

수식 셀은 openpyxl 저장 시 계산값이 남지 않아 파이썬이 못 읽는다.
그래서 판정·단일 테마 %는 여기서 다시 계산하고, FLAG는 값으로 입력받는다.
"""
import re, math, datetime

SCHEMA_CELL = "W1"
SCHEMA_V2 = "schema=v2"

ROWS = range(8, 15)
ROW_LABEL = {8: "펀더멘털·글로벌", 9: "펀더멘털·국내", 10: "통화정책·글로벌", 11: "통화정책·국내",
             12: "수급·글로벌TP", 13: "수급·국내", 14: "기술적·포지션"}
# 뉴스 합산으로 채점하면 안 되는 행: 정책은 '프라이싱 대비 갭', TP는 ACM z-score로 채점한다.
NO_NEWS_SCORE_ROWS = frozenset({10, 11, 12})

COL = {"note": 4, "dur_w": 5, "dur_prev": 6, "dur": 7, "cur_w": 8, "cur_prev": 9, "cur": 10,
       "v1_dur_w": 11, "v1_cur_w": 12, "dur_tag": 13, "cur_tag": 14,
       "news_dur": 15, "news_cur": 16, "digest": 17, "house": 18,
       "dur_contrib": 19, "cur_contrib": 20, "dur_share": 21, "cur_share": 22}
HUMAN_COLS = frozenset({COL["note"], COL["dur"], COL["cur"], COL["dur_tag"], COL["cur_tag"]})
AUTO_COLS = (COL["news_dur"], COL["news_cur"], COL["digest"], COL["house"])

DUR_WEIGHTS = {8: 0.20, 9: 0.10, 10: 0.20, 11: 0.20, 12: 0.10, 13: 0.15, 14: 0.05}
CUR_WEIGHTS = {8: 0.10, 9: 0.05, 10: 0.15, 11: 0.25, 12: 0.20, 13: 0.15, 14: 0.10}

FLAG = {"regime": "C19", "rho60": "C20", "transition": "C21", "event": "C22",
        "tp_break": "C23", "core_cpi_lt3": "C24", "threshold": "C25", "confirmed": "C26"}
MEMO_ROWS = range(28, 35)

# --- 점수 초안 (2026-09-21 사용자 결정: 자동 초안 → 사람이 조정·확정) ---
# 자동화는 사람 칸(D·G·J·M·N)에 초안을 쓰되, 아래 장부로 "사람이 고친 칸"을 가려 다시는 덮지 않는다.
DRAFT_CELL, DRAFT_MARK = "W2", "draft=1"
SHADOW = {"dur": 24, "cur": 25, "dur_tag": 26, "cur_tag": 27, "note": 28}   # X~AB: 자동화가 마지막으로 쓴 값
LOCK_COL = 29                                                                # AC: 사람이 고친 칸 목록 "dur,note"
DRAFT_FIELDS = ("dur", "cur", "dur_tag", "cur_tag", "note")

# --- 보고서형 서식 (2026-09-21 사용자 요청: 칸을 넘기지 말 것, 개조식, 회의에 쓸 수 있는 문체) ---
# D열(너비 ≈130, 14pt)은 한 줄에 한글 약 50자. 불릿은 한 줄에 끝나야 한눈에 읽힌다.
BULLET = "• "
NOTE_BULLETS, NOTE_BULLET_LEN = 3, 48        # D8:D14 — 불릿 최대 3개, 불릿당 48자
CONC_BULLETS, CONC_BULLET_LEN = 3, 80        # D3(D~J 병합, 한 줄 약 90자) — 머리말 1줄 + 불릿 3개
CONC_HEAD, CONC_HEAD_LEN = "【금주 결론】", 90
ROW_HEIGHT, CONC_ROW_HEIGHT = 84, 96         # 14pt 4줄 / 결론 4줄+여백
COMPACT_CELL, COMPACT_MARK = "W3", "compact=1"
_MARKERS = "*•·-–—▪■◦○●▶▷※"

TH_BASE, TH_EVENT, TH_STRONG = 0.5, 0.75, 1.0
NEWS_CAP, CURVE_CAP = 2.0, 1.5
DIVERGE_GAP = 1.0            # |내 점수 − 뉴스 참고| 가 이 이상이면 괴리로 표시
THEME_ALERT = 0.5

REGION_SPLIT = frozenset({"펀더멘털", "통화정책", "수급"})
HOUSE_PAT = re.compile(r"증권|자산운용|투자신탁|리서치센터|애널리스트")
LABELS = {"dur": ("O/W", "U/W"), "cur": ("Flattening", "Steepening")}


class SchemaError(RuntimeError):
    """v2가 아닌 워크북에 v2 스크립트가 쓰려 할 때."""


def is_v2(ws):
    return (ws[SCHEMA_CELL].value == SCHEMA_V2 or ws["B26"].value == "채점 확정"
            or ws["X7"].value == "초안(듀)")


def has_draft_layout(ws):
    return ws[DRAFT_CELL].value == DRAFT_MARK


def policy_score(gap_bp):
    """4-3: 갭(시장 프라이싱 − 내 경로, bp) → 듀레이션 점수. 커브는 부호 반대."""
    for th, s in ((50, 2), (25, 1)):
        if gap_bp >= th:
            return s
        if gap_bp <= -th:
            return -s
    return 0


def tp_score(z):
    """8-4: z = TP10 4주 변화 ÷ σ → −sign(z)·min(2, floor|z|). 듀레이션·커브 같은 부호."""
    mag = min(2, math.floor(abs(z)))
    return 0 if mag == 0 else int(-math.copysign(mag, z))


def require_v2(ws, path=""):
    if not is_v2(ws):
        raise SchemaError(
            f"v1 레이아웃 워크북입니다: {path}\n"
            f'  python "_news-to-macro/scripts/migrate_v2.py" "<파일>" --new-week <그 주 월요일 YYYY-MM-DD> 로 먼저 변환하세요. '
            f"(행 의미가 달라 그대로 쓰면 다른 칸에 점수가 들어갑니다)")


def _clip(text, limit, label, warns):
    text = " ".join(str(text).split())
    if len(text) <= limit:
        return text
    warns.append(f"{label}: 길이 초과 {len(text)}자 → {limit}자로 자름 (더 짧게 다시 쓸 것)")
    return text[:limit - 1].rstrip() + "…"


def split_items(text):
    """여러 줄 텍스트 → 불릿 기호를 뗀 항목 목록. 빈 줄 제외."""
    items = []
    for line in str(text or "").splitlines():
        line = line.strip()
        # '-'·'*'는 뒤에 공백이 있을 때만 불릿 기호로 본다("-0.5bp…"의 부호를 지키기 위함)
        while line and line[0] in _MARKERS and (line[0] not in "-*" or line[1:2] in ("", " ")):
            line = line[1:].lstrip()
        if line:
            items.append(line)
    return items


def to_bullets(items, max_n, max_len, label="요약"):
    """항목 목록 → ('• a\\n• b', 경고). 개수·길이 예산을 강제한다(칸을 넘기지 않게)."""
    warns = []
    if len(items) > max_n:
        warns.append(f"{label}: 불릿 {max_n}개 초과({len(items)}개) → 앞의 {max_n}개만 사용")
    lines = [BULLET + _clip(x, max_len, label, warns) for x in items[:max_n]]
    return "\n".join(lines), warns


def format_conclusion(text):
    """D3 결론: 머리말 한 줄(【금주 결론】 판정) + 불릿 최대 3개."""
    lines = [x for x in str(text or "").splitlines() if x.strip()]
    warns = []
    head = lines[0].strip() if lines else ""
    body = lines[1:]
    if not head.startswith(CONC_HEAD):
        if len(lines) > 1 or not head:
            head, body = CONC_HEAD, lines          # 머리말이 없으면 전부 불릿으로
        else:
            head = f"{CONC_HEAD} {head}"
    head = _clip(head, CONC_HEAD_LEN, "결론 머리말", warns)
    bullets, w = to_bullets(split_items("\n".join(body)), CONC_BULLETS, CONC_BULLET_LEN, "결론")
    return (head + ("\n" + bullets if bullets else "")), warns + w


def round_half(x):
    """0.5 단위 반올림, 부호 대칭(0.75→1.0, −0.75→−1.0)."""
    return math.copysign(math.floor(abs(x) * 2 + 0.5) / 2, x) if x else 0.0


def clamp(x, cap):
    return max(-cap, min(cap, x))


def map_row(factor, region):
    """News DB (요인, 지역) → Analysis 행. '기타'는 점수 행이 없다(이벤트 FLAG 후보)."""
    if factor == "기술적분석":
        return 14
    if factor not in REGION_SPLIT:
        return None
    glob = region == "글로벌"
    return {"펀더멘털": 8 if glob else 9, "통화정책": 10 if glob else 11,
            "수급": 12 if glob else 13}[factor]


def is_opinion(item):
    """애널리스트·하우스 의견인가. kind가 명시돼 있으면 그것이 우선, 없으면 출처·제목으로 판별."""
    kind = (item.get("kind") or "").strip()
    if kind:
        return kind == "의견"
    return bool(HOUSE_PAT.search(f"{item.get('source') or ''} {item.get('title') or ''}"))


def verdict(total, n_scored, block, event_on=False, relax=False):
    """6장 판정. block='dur'|'cur'. 7행이 다 채점되지 않았으면 판정하지 않는다(빈칸 금지)."""
    if n_scored < len(ROWS):
        return f"채점 미완 {n_scored}/{len(ROWS)}"
    hi, lo = LABELS[block]
    th = TH_EVENT if event_on else TH_BASE
    soften = relax and block == "dur"          # A′ + 근원 CPI<3% 완화는 U/W 쪽에만
    if total >= TH_STRONG:
        return f"강한 {hi}"
    if total >= th:
        return hi
    if total <= -TH_STRONG:
        return lo if soften else f"강한 {lo}"
    if total <= -th:
        return "Neutral" if soften else lo
    return "Neutral"


def theme_concentration(weights, scores, tags):
    """(최대 테마 비중, 태그). 기여 = |비중×점수|. 태그 없는 행은 분모에만 들어간다."""
    contrib = [abs((w or 0) * (s or 0)) for w, s in zip(weights, scores)]
    denom = sum(contrib)
    if denom == 0:
        return 0.0, None
    by_tag = {}
    for c, t in zip(contrib, tags):
        t = (t or "").strip()
        if t:
            by_tag = {**by_tag, t: by_tag.get(t, 0.0) + c}
    if not by_tag:
        return 0.0, None
    tag = max(by_tag, key=by_tag.get)
    return by_tag[tag] / denom, tag


def _num(v):
    try:
        return None if v is None or v == "" else float(v)
    except (TypeError, ValueError):
        return None


def _block(ws, w_col, s_col, tag_col, block, event_on, relax):
    weights = [_num(ws.cell(r, w_col).value) or 0.0 for r in ROWS]
    scores = [_num(ws.cell(r, s_col).value) for r in ROWS]
    tags = [ws.cell(r, tag_col).value for r in ROWS]
    n = sum(1 for s in scores if s is not None)
    total = round(sum(w * (s or 0.0) for w, s in zip(weights, scores)), 2)
    share, tag = theme_concentration(weights, scores, tags)
    return {"total": total, "n_scored": n, "verdict": verdict(total, n, block, event_on, relax),
            "theme": (round(share, 4), tag)}


def read_state(wb_path):
    """사람이 쓴 점수·태그·FLAG와 자동 참고 열을 읽어 판정까지 계산한다."""
    import openpyxl
    wb = openpyxl.load_workbook(wb_path, read_only=True)
    ws = wb["Analysis"]
    try:
        if not is_v2(ws):
            raise SchemaError(f"v1 레이아웃 워크북: {wb_path}")
        regime = str(ws[FLAG["regime"]].value or "").strip().replace("'", "′")
        flags = {"regime": regime or None,
                 "rho60": _num(ws[FLAG["rho60"]].value),
                 "transition": str(ws[FLAG["transition"]].value or "OFF").strip().upper(),
                 "event_on": str(ws[FLAG["event"]].value or "").strip().upper() == "ON",
                 "tp_break": str(ws[FLAG["tp_break"]].value or "OFF").strip().upper(),
                 "core_cpi_lt3": str(ws[FLAG["core_cpi_lt3"]].value or "").strip().upper() == "Y"}
        flags = {**flags, "relax": flags["regime"] == "A′" and flags["core_cpi_lt3"]}
        confirmed = str(ws[FLAG["confirmed"]].value or "").strip().upper() == "Y"
        locks = {r: {x for x in str(ws.cell(r, LOCK_COL).value or "").split(",") if x} for r in ROWS}
        rows = [{"row": r, "label": ROW_LABEL[r],
                 "draft_dur": _num(ws.cell(r, SHADOW["dur"]).value),
                 "draft_cur": _num(ws.cell(r, SHADOW["cur"]).value),
                 "locked": sorted(locks[r]),
                 "note": ws.cell(r, COL["note"]).value,
                 "dur": _num(ws.cell(r, COL["dur"]).value), "cur": _num(ws.cell(r, COL["cur"]).value),
                 "dur_prev": _num(ws.cell(r, COL["dur_prev"]).value),
                 "cur_prev": _num(ws.cell(r, COL["cur_prev"]).value),
                 "dur_tag": ws.cell(r, COL["dur_tag"]).value, "cur_tag": ws.cell(r, COL["cur_tag"]).value,
                 "news_dur": _num(ws.cell(r, COL["news_dur"]).value),
                 "news_cur": _num(ws.cell(r, COL["news_cur"]).value),
                 "house": ws.cell(r, COL["house"]).value} for r in ROWS]
        ev, rx = flags["event_on"], flags["relax"]
        return {"flags": flags, "rows": rows, "confirmed": confirmed,
                "edited_n": sum(len(v & {"dur", "cur"}) for v in locks.values()),
                "dur": _block(ws, COL["dur_w"], COL["dur"], COL["dur_tag"], "dur", ev, rx),
                "cur": _block(ws, COL["cur_w"], COL["cur"], COL["cur_tag"], "cur", ev, rx)}
    finally:
        wb.close()


def name_key(name):
    return "dur" if name == "듀" else "cur"


def divergences(state):
    """내 점수와 뉴스 참고 점수가 크게 다른 행 — 복기 포인트."""
    out = []
    for r in state["rows"]:
        for mine, ref, name in ((r["dur"], r["news_dur"], "듀"), (r["cur"], r["news_cur"], "커브")):
            if mine is not None and ref is not None and abs(mine - ref) >= DIVERGE_GAP:
                who = "내 점수" if name_key(name) in (r.get("locked") or ()) else "초안"
                out.append(f"{r['label']}({name}): {who} {mine:+g} vs 뉴스 참고 {ref:+g}")
    return out


def week_start(d):
    return d - datetime.timedelta(days=d.weekday())
