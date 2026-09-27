#!/usr/bin/env python3
"""점수 초안 기입 — 엑셀을 열면 요약·태그·점수 초안이 이미 채워져 있고, 사람은 점수만 조정·확정한다.

초안의 출처
  - 헤드리스 LLM이 쓴 draft.json (행별 요약·태그·점수 제안)
  - 없으면 롤업의 뉴스 참고 점수로 만든 자동 집계 초안(fallback) — LLM 인증이 풀려도 시트가 비지 않게

파이썬이 강제하는 규칙 (LLM 제안을 그대로 믿지 않는다)
  1. 점수는 −2~+2, 0.5 단위로 보정.
  2. **통화정책 2행은 방향으로 채점하지 않는다**(운영기준 4-3). draft에 market_pct·my_path_pct가 있으면
     갭(bp)으로 파이썬이 직접 채점(듀 = policy_score, 커브 = 반대 부호), 없으면 0 + "갭 수치 없음".
  3. 수급·글로벌 기간 프리미엄 행은 tp_4w_change_bp·sigma_bp가 있으면 z로 채점(8-4), adjust ±0.5,
     kw_agrees=false(Kim-Wright 부호 불일치)면 0. 없으면 0 + "ACM 수치 없음".
  4. 7행을 모두 채운다(빈칸 금지). 재료 없는 행은 0 + 사유.

보호 규칙 — **사람이 고친 칸은 그 주 동안 다시 덮어쓰지 않는다**
  각 칸(D·G·J·M·N)마다: 비어 있거나 '자동화가 마지막으로 쓴 값(X~AB 숨김 열)'과 같을 때만 갱신.
  다르면 사람이 고친 것 → AC열 잠금 목록에 올리고 영구 제외. (초안이 우연히 같은 값이 돼도 잠금 유지)
  채점 확정(C26=Y) 전까지 판정에는 "(초안)"이 붙는다.

사용법:  python draft_apply.py <workbook.xlsx> <draft.json>
"""
import sys, json
import openpyxl
from safe_io import safe_save
import sheet_v2 as sv
import migrate_v2 as mg

CELL_COL = {"dur": sv.COL["dur"], "cur": sv.COL["cur"], "dur_tag": sv.COL["dur_tag"],
            "cur_tag": sv.COL["cur_tag"], "note": sv.COL["note"]}
POLICY_ROWS, TP_ROW = (10, 11), 12
# 시트에 들어가는 문구는 보고서체로 짧게. 작업 안내("메모에 ~를 적으면…")는 시트가 아니라 알림 메일에 둔다.
NO_GAP = "갭 산출 불가(시장 프라이싱·정책경로 전망 미입력) → 0점"
NO_ACM = "ACM 기간 프리미엄 수치 미입력 → 0점"
NO_NEWS = "금주 관련 재료 없음 → 0점"


def num(v):
    try:
        return None if v is None or v == "" else float(v)
    except (TypeError, ValueError):
        return None


def coerce(v, label, warns):
    x = num(v)
    if x is None:
        return 0.0
    fixed = sv.round_half(sv.clamp(x, 2.0))
    if fixed != x:
        warns.append(f"{label}: {x:+g} → {fixed:+g} (−2~+2, 0.5 단위로 보정)")
    return fixed


def score_policy(row, r, warns):
    label = sv.ROW_LABEL[row]
    market, mine = num(r.get("market_pct")), num(r.get("my_path_pct"))
    if market is None or mine is None:
        if (num(r.get("dur")) or 0) or (num(r.get("cur")) or 0):
            warns.append(f"{label}: 갭 수치 없이 온 점수는 방향 채점이라 버림 → 0")
        return 0.0, 0.0, NO_GAP
    gap = round((market - mine) * 100)
    dur = float(sv.policy_score(gap))
    return dur, -dur, f"정책 갭 {gap:+d}bp (시장 {market:g}% − 전망 {mine:g}%)"


def score_tp(r, warns):
    chg, sigma = num(r.get("tp_4w_change_bp")), num(r.get("sigma_bp"))
    if chg is None or not sigma:
        if (num(r.get("dur")) or 0) or (num(r.get("cur")) or 0):
            warns.append(f"{sv.ROW_LABEL[TP_ROW]}: ACM 수치 없이 온 점수는 버림 → 0")
        return 0.0, 0.0, NO_ACM
    if r.get("kw_agrees") is False:
        return 0.0, 0.0, f"TP10 4주 {chg:+g}bp, Kim-Wright와 부호 불일치 → 0점"
    z = chg / sigma
    adj = sv.clamp(num(r.get("adjust")) or 0.0, 0.5)
    s = sv.round_half(sv.clamp(sv.tp_score(z) + adj, 2.0))
    return s, s, f"TP10 4주 {chg:+g}bp (z {z:+.2f})" + (f", 공급 조정 {adj:+g}" if adj else "")


def normalize(draft):
    """draft → ({row: {dur, cur, dur_tag, cur_tag, note}}, 경고목록). 7행을 모두 채워 돌려준다."""
    warns, given = [], {int(r["row"]): r for r in (draft or {}).get("rows", []) if num(r.get("row")) in set(sv.ROWS)}
    out = {}
    for row in sv.ROWS:
        r = given.get(row, {})
        items = sv.split_items(r.get("summary"))
        if row in POLICY_ROWS:
            dur, cur, basis = score_policy(row, r, warns)
            items = [basis] + items                  # 채점 근거가 항상 첫 불릿
        elif row == TP_ROW:
            dur, cur, basis = score_tp(r, warns)
            items = [basis] + items
        else:
            dur = coerce(r.get("dur"), f"{sv.ROW_LABEL[row]}(듀)", warns)
            cur = coerce(r.get("cur"), f"{sv.ROW_LABEL[row]}(커브)", warns)
            items = items or [NO_NEWS]
        summary, w = sv.to_bullets(items, sv.NOTE_BULLETS, sv.NOTE_BULLET_LEN, f"{sv.ROW_LABEL[row]} 요약")
        warns += w
        out[row] = {"dur": dur, "cur": cur, "note": summary,
                    "dur_tag": (str(r.get("dur_tag")).strip() or None) if r.get("dur_tag") else None,
                    "cur_tag": (str(r.get("cur_tag")).strip() or None) if r.get("cur_tag") else None,
                    "summary": summary}
    return out, warns


def fallback_draft(auto):
    """LLM 초안이 없을 때: 롤업의 뉴스 참고 점수를 그대로 초안으로. 정책·TP 행은 normalize가 0으로 만든다."""
    rows = []
    for f in (auto or {}).get("factors", []):
        digest = (f.get("digest") or "").split(": ", 1)[-1]          # "3건: a; b" → "a; b"
        items = ["자동 집계 초안(뉴스 강도 합산)"] + [x.strip() for x in digest.split(";") if x.strip()]
        rows.append({"row": f["row"], "dur": f.get("news_dur") or 0.0, "cur": f.get("news_curve") or 0.0,
                     "summary": "\n".join(items)})
    return {"rows": rows}


def same(a, b):
    if a is None or a == "":
        return b is None or b == ""
    na, nb = num(a), num(b)
    if na is not None and nb is not None and not isinstance(a, str) and not isinstance(b, str):
        return abs(na - nb) < 1e-9
    return str(a).strip() == str(b).strip()


def apply_draft(wb_path, draft, only_blank=False):
    """초안을 기입한다. only_blank=True면 빈 칸만 채운다(fallback이 LLM 초안을 밀어내지 않게)."""
    rows, warns = normalize(draft)
    wb = openpyxl.load_workbook(wb_path)
    ws = wb["Analysis"]
    sv.require_v2(ws, wb_path)
    if not sv.has_draft_layout(ws):
        mg.add_draft_layout(ws)
    if ws[sv.COMPACT_CELL].value != sv.COMPACT_MARK:
        mg.apply_compact_layout(ws)                # 1회만 — 이후 사용자가 바꾼 행 높이는 되돌리지 않는다
    written = locked = 0
    for row, vals in rows.items():
        locks = {x for x in str(ws.cell(row, sv.LOCK_COL).value or "").split(",") if x}
        for field in sv.DRAFT_FIELDS:
            cell, shadow = ws.cell(row, CELL_COL[field]), ws.cell(row, sv.SHADOW[field])
            blank = cell.value is None or cell.value == ""
            if field not in locks and not blank and not same(cell.value, shadow.value):
                locks.add(field)                      # 사람이 고친 칸 발견 → 이번 주 영구 제외
            if field in locks or (only_blank and not blank):
                continue
            cell.value = shadow.value = vals[field]
            written += 1
        ws.cell(row, sv.LOCK_COL).value = ",".join(sorted(locks)) or None
        locked += len(locks)
    safe_save(wb, wb_path, expect_sheets=["News DB", "Analysis"])
    for w in warns:
        print("  초안 보정: " + w)
    print(f"초안 기입: {written}칸 갱신, 사람이 고친 {locked}칸은 보존" + (" (빈 칸만)" if only_blank else ""))
    return {"written": written, "locked": locked, "warns": warns}


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__); sys.exit(1)
    with open(sys.argv[2], encoding="utf-8") as fp:
        apply_draft(sys.argv[1], json.load(fp))
