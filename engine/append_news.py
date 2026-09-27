#!/usr/bin/env python3
"""News DB 자동 추가 스크립트.

분류된 뉴스 JSON을 받아 매크로 분석 워크북의 'News DB' 시트(NewsDB 표)에
행을 추가하고, 표 범위를 확장한다. 집계(COUNTIF/AVERAGEIF 등)는 수식이라
엑셀에서 자동 재계산된다.

v2: 10번째 열 '유형'(지표/이벤트/의견)을 함께 기록한다. items에 kind가 없으면
출처·제목으로 판별한 값을 적어 넣어, 잘못 분류된 것이 엑셀에서 바로 보이게 한다.
v1 레이아웃 워크북에는 쓰기를 거부한다.

사용법:
  python append_news.py <workbook.xlsx> <items.json> [output.xlsx]

items.json 형식: references/classification_guide.md 참고
"""
import sys, json, datetime
from copy import copy
import openpyxl
from safe_io import safe_save
import sheet_v2 as sv

FACTORS = {"펀더멘털", "통화정책", "수급", "기술적분석", "기타"}
REGIONS = {"글로벌", "국내"}
DIRECTIONS = {"금리상방", "금리하방", "중립"}
KINDS = {"지표", "이벤트", "의견"}


def parse_date(v):
    if not v:
        return None
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d", "%m/%d", "%m.%d"):
        try:
            d = datetime.datetime.strptime(str(v).strip(), fmt)
            if fmt in ("%m/%d", "%m.%d"):
                d = d.replace(year=datetime.date.today().year)
            return d
        except ValueError:
            continue
    return str(v)  # 파싱 실패 시 원문 유지


def resolve_kind(item):
    kind = (item.get("kind") or "").strip()
    if kind in KINDS:
        return kind
    return "의견" if sv.is_opinion(item) else "지표"


def validate(item, idx):
    warns = []
    if item.get("factor") not in FACTORS:
        warns.append(f"[{idx}] 요인 '{item.get('factor')}' 비표준")
    if item.get("region") not in REGIONS:
        warns.append(f"[{idx}] 지역 '{item.get('region')}' 비표준")
    if item.get("direction") not in DIRECTIONS:
        warns.append(f"[{idx}] 방향성 '{item.get('direction')}' 비표준")
    if item.get("kind") and item.get("kind") not in KINDS:
        warns.append(f"[{idx}] 유형 '{item.get('kind')}' 비표준 → 자동 판별값 사용")
    if resolve_kind(item) == "의견" and not (item.get("source") or "").strip():
        warns.append(f"[{idx}] 의견인데 출처(하우스명) 없음 — 하우스 뷰 집계에 '미상'으로 들어감")
    d, s = item.get("direction"), item.get("intensity")
    if isinstance(s, (int, float)):
        if d == "금리상방" and s < 0: warns.append(f"[{idx}] 금리상방인데 강도 음수({s})")
        if d == "금리하방" and s > 0: warns.append(f"[{idx}] 금리하방인데 강도 양수({s})")
    return warns


def append(wb_path, json_path, out_path=None):
    with open(json_path, encoding="utf-8") as fp:
        items = json.load(fp).get("items", [])
    if not items:
        print("추가할 뉴스가 없습니다."); return 0

    wb = openpyxl.load_workbook(wb_path)
    sv.require_v2(wb["Analysis"], wb_path)
    ws = wb["News DB"]
    tbl = ws.tables["NewsDB"]
    start, end = tbl.ref.split(":")          # 예) A4:J26
    end_col = "".join(c for c in end if c.isalpha())
    end_row = int("".join(c for c in end if c.isdigit()))
    tmpl_row = end_row                       # 스타일 복사용 템플릿(마지막 데이터 행)

    all_warns = []
    r = end_row
    for i, it in enumerate(items, 1):
        all_warns += validate(it, i)
        r += 1
        vals = [parse_date(it.get("date")), it.get("factor", ""), it.get("region", ""),
                it.get("keyword", ""), it.get("title", ""), it.get("content", ""),
                it.get("direction", ""), it.get("intensity"), it.get("source", ""), resolve_kind(it)]
        for c_idx, val in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c_idx, value=val)
            src = ws.cell(row=tmpl_row, column=c_idx)
            if src.has_style:
                cell.font = copy(src.font)
                cell.border = copy(src.border)
                cell.fill = copy(src.fill)
                cell.alignment = copy(src.alignment)
                cell.number_format = src.number_format
            if c_idx == 1 and isinstance(vals[0], (datetime.date, datetime.datetime)):
                cell.number_format = "yyyy-mm-dd"

    tbl.ref = f"{start}:{end_col}{r}"
    if tbl.autoFilter is not None:
        tbl.autoFilter.ref = tbl.ref
    safe_save(wb, out_path or wb_path, expect_sheets=["News DB", "Analysis"])

    print(f"뉴스 {len(items)}건 추가 (행 {end_row+1}~{r}). 표 범위 → {tbl.ref}")
    print(f"저장: {out_path or wb_path}")
    if all_warns:
        print("\n검토 필요:")
        for w in all_warns:
            print("  -", w)
    return len(items)


def main():
    if len(sys.argv) < 3:
        print(__doc__); sys.exit(1)
    try:
        append(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None)
    except sv.SchemaError as e:
        print("실패: " + str(e), file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
